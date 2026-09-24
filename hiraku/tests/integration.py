"""Opt-in GPU acceptance checks using generated documents only."""
from __future__ import annotations
import json
import sys
import time
from pathlib import Path

import msoffcrypto
import pyzipper
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pypdf import PdfReader

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from app import Library
from test_core import make_pdf

ROOT=Path(__file__).resolve().parents[1] / '.test-data' / 'integration'


def wait(library,jid,timeout=100):
    deadline=time.monotonic()+timeout
    while library.jobs[jid]['state'] in {'queued','recovering','pausing','converting'}:
        if time.monotonic()>deadline:
            library.stop(jid,cancel=True)
            raise TimeoutError(jid)
        time.sleep(.25)
    state=library.jobs[jid]['state']
    if state!='ready':
        raise AssertionError(json.dumps(library.jobs[jid],ensure_ascii=False))


def main():
    ROOT.mkdir(parents=True,exist_ok=True)
    library=Library(ROOT/'library')
    results=[]
    try:
        sources=[]
        for ext in ('.docx','.xlsx','.pptx'):
            plain=ROOT/('plain'+ext)
            if ext=='.docx':
                doc=Document();doc.add_paragraph('Hiraku acceptance test');doc.save(plain)
            elif ext=='.xlsx':
                doc=Workbook();doc.active['A1']='Hiraku';doc.active['B1']='=2+3';doc.save(plain)
            else:
                doc=Presentation();slide=doc.slides.add_slide(doc.slide_layouts[0]);slide.shapes.title.text='Hiraku';doc.save(plain)
            encrypted=ROOT/('encrypted'+ext)
            with plain.open('rb') as src,encrypted.open('wb') as dest:
                msoffcrypto.OfficeFile(src).encrypt('Hiraku42',dest)
            sources.append(encrypted)
        encrypted=ROOT/'encrypted.zip'
        with pyzipper.AESZipFile(str(encrypted),'w',encryption=pyzipper.WZ_AES,compression=pyzipper.ZIP_DEFLATED) as z:
            z.setpassword(b'Hiraku42');z.writestr('folder/acceptance.txt',b'Original content\n'*10)
        sources.append(encrypted)
        for algorithm in ('RC4-40', 'RC4-128', 'AES-128', 'AES-256-R5', 'AES-256'):
            encrypted=ROOT/(algorithm+'.pdf');make_pdf(encrypted,'Hiraku42',algorithm);sources.append(encrypted)
        for source in sources:
            jid=library.import_pdf(source.name,source.read_bytes())
            started=time.monotonic()
            if not library.jobs[jid].get('unlocked'):
                library.start_recovery(jid,{'strategy':'dictionary','words':'wrong\nHiraku42\nother','devices':'1','minutes':1})
                wait(library,jid)
            result={'format':source.name,'state':library.jobs[jid]['state'],'seconds':round(time.monotonic()-started,2)}
            results.append(result);print(json.dumps(result),flush=True)
        (ROOT/'results.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
    finally:
        library.close()


if __name__=='__main__':main()

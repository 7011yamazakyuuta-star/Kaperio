from __future__ import annotations

import io
import json
import tempfile
import threading
import time
import unittest
import zipfile
from pathlib import Path

import msoffcrypto
import pyzipper
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas

from app import Library, acquire_instance
from formats import get_hash, inspect_file, unlock_file
from pdf_tools import export_pdf, extract_hash, preview_png
from recovery import parse_status, validate_plan, write_inputs


def make_pdf(path, password=None, algorithm='RC4-40'):
    stream = io.BytesIO()
    c = canvas.Canvas(stream, pagesize=(420, 595))
    for i in range(2):
        c.drawString(30, 550, 'Kaperio test page ' + str(i + 1))
        c.rect(30, 30, 360, 480)
        c.showPage()
    c.save()
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(stream.getvalue())))
    if password is not None:
        writer.encrypt(password, algorithm=algorithm)
    writer.write(path)


class CoreTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_pdf_revisions_and_verification(self):
        from formats import unlock_file
        for algorithm, mode in [('RC4-40',10400),('RC4-128',10500),('AES-128',10500),('AES-256-R5',10600),('AES-256',10700)]:
            source = self.root / 'test.pdf'
            make_pdf(source, 'Test42', algorithm)
            info = inspect_file(source, '.pdf')
            value, actual = extract_hash(PdfReader(source))
            self.assertEqual(mode, actual)
            self.assertTrue(value.startswith('$pdf$'))
            self.assertEqual(len(value[5:].split('*')), 15 if mode in (10600,10700) else 11)
            target = self.root / 'unlocked.pdf'
            with self.assertRaises(ValueError):
                unlock_file(source, 'wrong', target, info)
            self.assertEqual(unlock_file(source, 'Test42', target, info), 2)
            self.assertFalse(PdfReader(target).is_encrypted)
            self.assertIn('test page 1', PdfReader(target).pages[0].extract_text())

    def test_pdf_exports(self):
        source = self.root / 'source.pdf'
        make_pdf(source)
        self.assertTrue(preview_png(source, 0).startswith(b'\x89PNG'))
        for kind, suffix in [('image_pdf','.pdf'),('images','.zip'),('word','.docx'),('text','.md')]:
            target = self.root / ('out'+suffix)
            export_pdf(source,target,kind,100,False,lambda *_:None,lambda:False)
            self.assertTrue(target.exists())
            if kind == 'image_pdf':
                self.assertEqual(len(PdfReader(target).pages),2)
            elif kind == 'images':
                with zipfile.ZipFile(target) as z:
                    self.assertEqual(len(z.namelist()),2)
            elif kind == 'word':
                self.assertEqual(len(Document(target).inline_shapes),2)
        target = self.root / 'cancel.pdf'
        with self.assertRaises(InterruptedError):
            export_pdf(source,target,'image_pdf',100,False,lambda *_:None,lambda:True)
        self.assertFalse(target.exists())

    def test_office_decryption_roundtrip(self):
        for extension in ('.docx','.xlsx','.pptx'):
            source = self.root / ('plain'+extension)
            if extension == '.docx':
                doc=Document();doc.add_paragraph('Kaperio');doc.save(source)
            elif extension == '.xlsx':
                doc=Workbook();doc.active['A1']='Kaperio';doc.active['B1']='=1+2';doc.save(source)
            else:
                doc=Presentation();slide=doc.slides.add_slide(doc.slide_layouts[0]);slide.shapes.title.text='Kaperio';doc.save(source)
            encrypted = self.root / ('encrypted'+extension)
            with source.open('rb') as src, encrypted.open('wb') as dest:
                msoffcrypto.OfficeFile(src).encrypt('Test42',dest)
            info=inspect_file(encrypted,extension)
            self.assertEqual(info['mode'],9600)
            target=self.root / ('unlocked'+extension)
            with self.assertRaises(ValueError):
                unlock_file(encrypted,'wrong',target,info)
            unlock_file(encrypted,'Test42',target,info)
            self.assertEqual(source.read_bytes(),target.read_bytes())

    def test_zip_roundtrip_and_traversal(self):
        source=self.root/'test.zip'
        with pyzipper.AESZipFile(str(source),'w',encryption=pyzipper.WZ_AES) as z:
            z.setpassword(b'Test42');z.writestr('folder/test.txt',b'content preserved')
        info=inspect_file(source,'.zip')
        self.assertTrue(info['encrypted'])
        target=self.root/'unlocked.zip'
        unlock_file(source,'Test42',target,info)
        with zipfile.ZipFile(target) as z:
            self.assertEqual(z.read('folder/test.txt'),b'content preserved')
            self.assertFalse(z.infolist()[0].flag_bits & 1)
        with zipfile.ZipFile(source,'w') as z:z.writestr('../outside.txt','bad')
        with self.assertRaises(ValueError):inspect_file(source,'.zip')

    def test_masks_keep_suffix_and_count(self):
        plan=validate_plan({'min':4,'max':5,'prefix':'A','suffix':'9','charsets':['lower']})
        self.assertEqual(plan['candidates'],str(26**2+26**3))
        write_inputs(self.root,plan)
        self.assertEqual((self.root/'search.hcmask').read_text().splitlines(),['?l,A?1?19','?l,A?1?1?19'])
        with self.assertRaises(ValueError):validate_plan({'prefix':'abc','min':1})
        with self.assertRaises(ValueError):validate_plan({'devices':'1 --force'})

    def test_status_redacts_candidates(self):
        line=json.dumps({'status':3,'progress':[10,100],'devices':[{'speed':12,'temp':50,'candidates':'secret'}]})
        status=parse_status(line)
        self.assertEqual(status['speed'],12)
        self.assertNotIn('secret',json.dumps(status))

    def test_single_instance_and_removal(self):
        lock=acquire_instance(self.root)
        self.assertIsNotNone(lock)
        self.assertIsNone(acquire_instance(self.root))
        lock.close()
        lock=acquire_instance(self.root)
        self.assertIsNotNone(lock)
        lock.close()
        source=self.root/'source.pdf';make_pdf(source)
        library=Library(self.root/'library')
        try:
            jid=library.import_pdf(source.name,source.read_bytes())
            library.remove(jid)
            self.assertEqual(library.snapshot(),[])
            self.assertTrue(source.exists())
        finally:library.close()

    def test_library_persistence_and_wrong_password(self):
        source=self.root/'test.pdf';make_pdf(source,'Test42')
        data=self.root/'data';library=Library(data)
        try:
            jid=library.import_pdf('日本語.pdf',source.read_bytes())
            self.assertEqual(library.import_pdf('duplicate.pdf',source.read_bytes()),jid)
            with self.assertRaises(ValueError):library.known_password(jid,'wrong')
            self.assertEqual(library.jobs[jid]['state'],'locked')
            library.known_password(jid,'Test42')
            self.assertNotIn('Test42',(data/jid/'job.json').read_text(encoding='utf-8'))
            restored=Library(data)
            try:self.assertTrue(restored.snapshot()[0]['available'])
            finally:restored.close()
        finally:library.close()


if __name__=='__main__':
    (Path(__file__).resolve().parents[1]/'.test-data').mkdir(exist_ok=True)
    unittest.main()

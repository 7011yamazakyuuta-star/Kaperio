"""Opt-in synthetic large-document benchmark. Requires psutil; never uses user files."""
import argparse
import hashlib
import http.client
import json
import os
import platform
import random
import subprocess
import tempfile
import threading
import time
import zipfile
from pathlib import Path
from urllib.parse import parse_qs, quote, urlsplit

import msoffcrypto
import psutil
import pyzipper
from PIL import Image
from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.drawing.image import Image as SheetImage
from pptx import Presentation
from pptx.util import Inches as SlideInches
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DictionaryObject, NameObject, NumberObject, DecodedStreamObject

PASSWORD = 'SyntheticLarge42'
MIB = 1024 * 1024


def digest(path):
    with Path(path).open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def prepare(root):
    root.mkdir(parents=True, exist_ok=True)
    manifest = root / 'fixtures.json'
    if manifest.exists():
        return json.loads(manifest.read_text())
    rng = random.Random(20260925)
    images = []
    for index in range(16):
        image = Image.frombytes('RGB', (2048, 2048), rng.randbytes(2048 * 2048 * 3))
        path = root / f'image-{index:02d}.png'
        image.save(path)
        image.close()
        images.append(path)
    records = []
    writer = PdfWriter()
    for path in images:
        page = writer.add_blank_page(width=420, height=595)
        with Image.open(path) as image:
            picture = DecodedStreamObject()
            picture.set_data(image.tobytes())
        picture.update({NameObject('/Type'): NameObject('/XObject'), NameObject('/Subtype'): NameObject('/Image'),
                        NameObject('/Width'): NumberObject(2048), NameObject('/Height'): NumberObject(2048),
                        NameObject('/ColorSpace'): NameObject('/DeviceRGB'), NameObject('/BitsPerComponent'): NumberObject(8)})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/XObject'): DictionaryObject({NameObject('/Im0'): writer._add_object(picture)})})
        content = DecodedStreamObject()
        content.set_data(b'q 420 0 0 595 0 0 cm /Im0 Do Q')
        page[NameObject('/Contents')] = writer._add_object(content)
    writer.encrypt(PASSWORD, algorithm='AES-256')
    pdf = root / 'large.pdf'
    writer.write(pdf)
    writer.close()
    records.append({'name': pdf.name, 'pages': len(images)})
    del writer
    for extension in ('docx', 'xlsx', 'pptx'):
        plain = root / ('plain.' + extension)
        if extension == 'docx':
            document = Document()
            for index, path in enumerate(images):
                document.add_paragraph(f'Synthetic image {index + 1}')
                document.add_picture(str(path), width=Inches(5))
                if index < len(images) - 1:
                    document.add_page_break()
            document.save(plain)
        elif extension == 'xlsx':
            document = Workbook()
            for index, path in enumerate(images):
                sheet = document.active if index == 0 else document.create_sheet()
                sheet.title = f'Image {index + 1}'
                sheet['A1'] = f'Synthetic image {index + 1}'
                picture = SheetImage(str(path))
                picture.width = picture.height = 600
                sheet.add_image(picture, 'A3')
            document.save(plain)
        else:
            document = Presentation()
            for path in images:
                slide = document.slides.add_slide(document.slide_layouts[6])
                slide.shapes.add_picture(str(path), SlideInches(1), SlideInches(.5), width=SlideInches(6))
            document.save(plain)
        target = root / ('large.' + extension)
        with plain.open('rb') as source, target.open('wb') as output:
            msoffcrypto.OfficeFile(source).encrypt(PASSWORD, output)
        records.append({'name': target.name, 'plain_sha256': digest(plain)})
        del document
    target = root / 'large.zip'
    with pyzipper.AESZipFile(target, 'w', compression=zipfile.ZIP_STORED, encryption=pyzipper.WZ_AES) as archive:
        archive.setpassword(PASSWORD.encode())
        for path in images:
            archive.write(path, path.name)
    records.append({'name': target.name, 'entries': {path.name: digest(path) for path in images}})
    for record in records:
        path = root / record['name']
        record.update(bytes=path.stat().st_size, sha256=digest(path))
        assert 190 * MIB < record['bytes'] <= 200 * MIB, record
    manifest.write_text(json.dumps(records, indent=2), encoding='utf-8')
    return records


class Measurements:
    def __init__(self, pid):
        self.process = psutil.Process(pid)
        self.stage = 'startup'
        self.values = {}
        self.stop = threading.Event()
        self.thread = threading.Thread(target=self.sample, daemon=True)
        self.thread.start()

    def sample(self):
        while not self.stop.is_set():
            try:
                processes = [self.process] + self.process.children(recursive=True)
                rss = private = 0
                for process in processes:
                    try:
                        memory = process.memory_info()
                        rss += memory.rss
                        private += getattr(memory, 'private', 0)
                    except psutil.Error:
                        pass
                value = self.values.setdefault(self.stage, {'peak_rss_bytes': 0, 'peak_private_bytes': 0})
                value['peak_rss_bytes'] = max(value['peak_rss_bytes'], rss)
                value['peak_private_bytes'] = max(value['peak_private_bytes'], private)
            except psutil.Error:
                pass
            self.stop.wait(.02)

    def run(self, stage, operation):
        self.stage = stage
        started = time.perf_counter()
        try:
            return operation()
        finally:
            self.values.setdefault(stage, {})['seconds'] = round(time.perf_counter() - started, 3)


def run_case(executable, fixtures, record, work):
    with tempfile.TemporaryDirectory(prefix='large-run-', dir=work) as temporary:
        root = Path(temporary)
        process = subprocess.Popen([str(executable), '--port', '0', '--data', str(root), '--no-browser'],
                                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                                   creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
        measurements = Measurements(process.pid)
        url = token = None
        try:
            for _ in range(300):
                if process.poll() is not None:
                    raise RuntimeError('App exited during startup')
                try:
                    url = urlsplit(json.loads((root / 'launch.json').read_text())['url'])
                    token = parse_qs(url.query)['token'][0]
                    break
                except (FileNotFoundError, ValueError):
                    time.sleep(.1)
            if url is None:
                raise TimeoutError('Startup')

            def request(route, body=None, upload=None, download=None):
                client = http.client.HTTPConnection(url.hostname, url.port, timeout=600)
                headers = {'Cookie': 'loxmit_session=' + token, 'X-Loxmit': '1'}
                source = None
                try:
                    if upload:
                        headers.update({'X-Filename': quote(upload.name), 'Content-Length': str(upload.stat().st_size)})
                        source = upload.open('rb')
                        body = source
                    elif body is not None:
                        body = json.dumps(body).encode()
                    client.request('GET' if body is None else 'POST', route, body, headers)
                    response = client.getresponse()
                    if response.status != 200:
                        raise AssertionError((route, response.status, response.read(2000)))
                    if download:
                        sha = hashlib.sha256()
                        with download.open('wb') as output:
                            while chunk := response.read(MIB):
                                output.write(chunk)
                                sha.update(chunk)
                        return sha.hexdigest()
                    return json.loads(response.read())
                finally:
                    if source:
                        source.close()
                    client.close()

            version = request('/api/settings')['version']
            assert request('/api/jobs')['jobs'] == []
            source = fixtures / record['name']
            assert digest(source) == record['sha256']
            jid = measurements.run('import', lambda: request('/api/import', upload=source)['id'])
            measurements.run('unlock', lambda: request('/api/jobs/' + jid + '/unlock', {'password': PASSWORD}))
            job = next(j for j in request('/api/jobs')['jobs'] if j['id'] == jid)
            output = next(o for o in job['outputs'] if o['kind'] == 'unlocked')
            target = root / ('download' + source.suffix)
            actual = measurements.run('download', lambda: request('/api/jobs/' + jid + '/download/' + output['file'], download=target))
            measurements.stage = 'verification'
            if 'plain_sha256' in record:
                assert actual == record['plain_sha256']
            elif source.suffix == '.zip':
                with zipfile.ZipFile(target) as archive:
                    assert set(archive.namelist()) == set(record['entries'])
                    for name, expected in record['entries'].items():
                        with archive.open(name) as stream:
                            assert hashlib.file_digest(stream, 'sha256').hexdigest() == expected
            else:
                with source.open('rb') as encrypted, target.open('rb') as decrypted:
                    original, result = PdfReader(encrypted), PdfReader(decrypted)
                    assert original.decrypt(PASSWORD) and not result.is_encrypted
                    assert len(result.pages) == record['pages']
                    for old, new in zip(original.pages, result.pages):
                        assert hashlib.sha256(old['/Resources']['/XObject']['/Im0'].get_data()).digest() == hashlib.sha256(new['/Resources']['/XObject']['/Im0'].get_data()).digest()
                for kind in ('image_pdf', 'word', 'text'):
                    def export():
                        request('/api/jobs/' + jid + '/export', {'kind': kind, 'dpi': 100})
                        deadline = time.monotonic() + 600
                        while time.monotonic() < deadline:
                            current = next(j for j in request('/api/jobs')['jobs'] if j['id'] == jid)
                            if current['state'] not in ('converting', 'queued'):
                                assert current['state'] == 'ready', current
                                item = next(o for o in current['outputs'] if o['kind'] == kind)
                                return root / jid / item['file']
                            time.sleep(.1)
                        raise TimeoutError('Export ' + kind)
                    artifact = measurements.run('export_' + kind, export)
                    measurements.stage = 'verification'
                    if kind == 'image_pdf':
                        with artifact.open('rb') as stream:
                            assert len(PdfReader(stream).pages) == record['pages']
                        import pypdfium2 as pdfium
                        with pdfium.PdfDocument(artifact) as document:
                            page = document[0]
                            try:
                                bitmap = page.render(scale=.25)
                                image = bitmap.to_pil()
                                assert any(low != high for low, high in image.getextrema())
                                image.close()
                                bitmap.close()
                            finally:
                                page.close()
                    elif kind == 'word':
                        assert len(Document(artifact).inline_shapes) == record['pages']
                    else:
                        assert artifact.read_text(encoding='utf-8').count('## Page ') == record['pages']
            return {'file': record['name'], 'bytes': record['bytes'], 'sha256': record['sha256'],
                    'app_version': version, 'passed': True, 'stages': measurements.values}
        finally:
            measurements.stop.set()
            measurements.thread.join()
            if url and token and process.poll() is None:
                try:
                    request('/api/shutdown', {})
                except Exception:
                    pass
            try:
                process.wait(timeout=15)
            except subprocess.TimeoutExpired:
                process.terminate()
                process.wait(timeout=15)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--fixtures', type=Path, required=True)
    parser.add_argument('--prepare', action='store_true')
    parser.add_argument('--executable', type=Path)
    parser.add_argument('--report', type=Path)
    args = parser.parse_args()
    if args.prepare:
        print(json.dumps(prepare(args.fixtures), indent=2), flush=True)
        return
    records = json.loads((args.fixtures / 'fixtures.json').read_text())
    report = {'system': platform.platform(), 'physical_memory_bytes': psutil.virtual_memory().total,
              'sample_interval_ms': 20, 'executable_sha256': digest(args.executable), 'cases': [],
              'limitations': ['Synthetic image-heavy documents; not all real documents.', 'Known-password decryption, not GPU recovery.',
                              'RSS includes app descendants, excludes fixture generator, client and OS file cache.',
                              'Office-to-PDF conversion is not exercised. Sampled peaks can miss shorter spikes.']}
    args.report.parent.mkdir(parents=True, exist_ok=True)
    for record in records:
        print('Testing ' + record['name'], flush=True)
        result = run_case(args.executable.resolve(), args.fixtures.resolve(), record, args.report.parent)
        report['cases'].append(result)
        args.report.write_text(json.dumps(report, indent=2), encoding='utf-8')
        print(json.dumps(result), flush=True)


if __name__ == '__main__':
    main()

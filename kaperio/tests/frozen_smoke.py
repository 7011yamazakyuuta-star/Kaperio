"""Exercise the shipped executable over HTTP; no GPU, network, or private files."""
import argparse
import http.client
import io
import json
import os
import platform
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path
from urllib.parse import urlsplit, parse_qs, quote

import msoffcrypto
import pyzipper
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pypdf import PdfReader

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from test_core import make_pdf


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('executable', type=Path)
    parser.add_argument('--hashcat', type=Path)
    parser.add_argument('--zip2john', type=Path)
    args = parser.parse_args()
    executable = args.executable.resolve()
    with tempfile.TemporaryDirectory(prefix='kaperio-smoke-') as temporary:
        root = Path(temporary)
        crypto_report = root / 'crypto.json'
        crypto_check = subprocess.run([str(executable), '--self-check', str(crypto_report)],
                                      cwd=root, timeout=60, capture_output=True)
        assert crypto_report.exists(), ('No crypto self-check report', crypto_check.returncode)
        crypto_result = json.loads(crypto_report.read_text())
        assert crypto_check.returncode == 0 and crypto_result['ok'], crypto_result
        data = root / 'library'
        command = [str(executable), '--port', '0', '--data', str(data), '--no-browser']
        environment = dict(os.environ, KAPERIO_HASHCAT='', KAPERIO_ZIP2JOHN='')
        process = subprocess.Popen(command, cwd=root, env=environment, stdout=subprocess.DEVNULL,
                                   stderr=subprocess.DEVNULL)
        launch = None
        checks = []
        try:
            for _ in range(300):
                if process.poll() is not None:
                    raise RuntimeError('Executable exited before startup: ' + str(process.returncode))
                try:
                    launch = json.loads((data / 'launch.json').read_text())
                    break
                except (FileNotFoundError, ValueError):
                    time.sleep(.1)
            if not launch:
                raise TimeoutError('Executable did not start')
            url = urlsplit(launch['url'])
            token = parse_qs(url.query)['token'][0]

            def request(path, body=None, name=None, authenticated=True):
                client = http.client.HTTPConnection(url.hostname, url.port, timeout=30)
                headers = {'X-Kaperio': '1'}
                if authenticated:
                    headers['Cookie'] = 'kaperio_session=' + token
                if name:
                    headers['X-Filename'] = quote(name)
                if isinstance(body, dict):
                    body = json.dumps(body).encode()
                client.request('GET' if body is None else 'POST', path, body, headers)
                response = client.getresponse()
                result = response.read()
                status = response.status
                client.close()
                return status, result

            def api(path, body=None, name=None):
                status, raw = request(path, body, name)
                assert status == 200, (path, status, raw[:1000])
                return json.loads(raw)

            def job(jid):
                return next(j for j in api('/api/jobs')['jobs'] if j['id'] == jid)

            def wait_export(jid):
                for _ in range(300):
                    value = job(jid)
                    if value['state'] not in ('converting', 'queued', 'recovering'):
                        assert value['state'] == 'ready', value['message']
                        return value
                    time.sleep(.1)
                raise TimeoutError('Export did not complete')

            assert request('/api/jobs', authenticated=False)[0] == 403
            assert b'Kaperio' in request('/')[1]
            notices = request('/api/licenses')
            assert notices[0] == 200 and b'Python' in notices[1] and b'pdfium' in notices[1], (notices[0], notices[1][:500])
            assert api('/api/settings')['version'] == '0.2.0-alpha.1'
            if args.hashcat:
                api('/api/settings', {'hashcat': str(args.hashcat.resolve()),
                                     'zip2john': str(args.zip2john.resolve()) if args.zip2john else ''})
            checks += ['crypto-self-check', 'auth', 'static', 'bundled-licenses', 'version']
            pdf = root / 'sample.pdf'
            make_pdf(pdf, 'Test42', 'AES-256')
            jid = api('/api/import', pdf.read_bytes(), pdf.name)['id']
            assert request('/api/jobs/' + jid + '/unlock', {'password': 'wrong'})[0] == 400
            api('/api/jobs/' + jid + '/unlock', {'password': 'Test42'})
            unlocked = data / jid / 'unlocked.pdf'
            assert not PdfReader(unlocked).is_encrypted
            preview = request('/api/jobs/' + jid + '/preview')
            assert preview[0] == 200 and preview[1].startswith(b'\x89PNG')
            for kind in ('image_pdf', 'word', 'images', 'text'):
                api('/api/jobs/' + jid + '/export', {'kind': kind, 'dpi': 100})
                value = wait_export(jid)
                output = next(o for o in value['outputs'] if o['kind'] == kind)
                assert request('/api/jobs/' + jid + '/download/' + output['file'])[0] == 200
            checks += ['pdf-aes256-unlock', 'pdfium-preview', 'four-export-formats']
            for algorithm in ('RC4-40', 'AES-128', 'AES-256-R5'):
                variant = root / (algorithm + '.pdf')
                make_pdf(variant, 'Test42', algorithm)
                vid = api('/api/import', variant.read_bytes(), variant.name)['id']
                api('/api/jobs/' + vid + '/unlock', {'password': 'Test42'})
                assert not PdfReader(data / vid / 'unlocked.pdf').is_encrypted
                checks.append('frozen-pdf-' + algorithm)
            if args.hashcat:
                gpu_pdf = root / 'gpu.pdf'
                make_pdf(gpu_pdf, 'Test42')
                gid = api('/api/import', gpu_pdf.read_bytes(), gpu_pdf.name)['id']
                api('/api/jobs/' + gid + '/recover', {'strategy': 'dictionary', 'words': 'wrong\nTest42',
                                                    'devices': '1', 'minutes': 1})
                wait_export(gid)
                checks.append('frozen-hashcat-pdf')
            for extension in ('.docx', '.xlsx', '.pptx'):
                plain = root / ('plain' + extension)
                if extension == '.docx':
                    doc = Document(); doc.add_paragraph('Kaperio'); doc.save(plain)
                elif extension == '.xlsx':
                    doc = Workbook(); doc.active['A1'] = 'Kaperio'; doc.save(plain)
                else:
                    doc = Presentation(); doc.slides.add_slide(doc.slide_layouts[0]); doc.save(plain)
                encrypted = io.BytesIO()
                with plain.open('rb') as src:
                    msoffcrypto.OfficeFile(src).encrypt('Test42', encrypted)
                oid = api('/api/import', encrypted.getvalue(), 'encrypted' + extension)['id']
                assert job(oid)['info']['mode'] == 9600, job(oid)['info']
                if args.hashcat:
                    compound = {
                        '.docx': {'strategy': 'dictionary_rules', 'words': 'wrong\ntest4'},
                        '.xlsx': {'strategy': 'hybrid_suffix', 'words': 'wrong\nTest',
                                  'min': 2, 'max': 2, 'charsets': ['digits']},
                        '.pptx': {'strategy': 'hybrid_prefix', 'words': 'wrong\nest42',
                                  'min': 1, 'max': 1, 'charsets': ['upper']},
                    }[extension]
                    api('/api/jobs/' + oid + '/recover', dict(compound, devices='1', minutes=1))
                    wait_export(oid)
                    checks.append('frozen-hashcat-' + extension[1:])
                else:
                    api('/api/jobs/' + oid + '/unlock', {'password': 'Test42'})
                assert (data / oid / ('unlocked' + extension)).read_bytes() == plain.read_bytes()
                checks.append('office-helper-unlock-' + extension[1:])
            archive = io.BytesIO()
            with pyzipper.AESZipFile(archive, 'w', encryption=pyzipper.WZ_AES) as zipped:
                zipped.setpassword(b'Test42'); zipped.writestr('test.txt', b'preserved')
            zid = api('/api/import', archive.getvalue(), 'test.zip')['id']
            if args.hashcat and args.zip2john:
                api('/api/jobs/' + zid + '/recover', {'strategy': 'dictionary', 'words': 'wrong\nTest42',
                                                    'devices': '1', 'minutes': 1})
                wait_export(zid)
                checks.append('frozen-hashcat-zip')
            else:
                api('/api/jobs/' + zid + '/unlock', {'password': 'Test42'})
            with zipfile.ZipFile(data / zid / 'unlocked.zip') as zipped:
                assert zipped.read('test.txt') == b'preserved'
            checks.append('zip-aes-unlock')
            second = subprocess.run(command, cwd=root, env=environment, timeout=30, capture_output=True)
            assert second.returncode == 0
            for value in api('/api/jobs')['jobs']:
                api('/api/jobs/' + value['id'] + '/remove', {})
            assert pdf.exists() and not api('/api/jobs')['jobs']
            checks += ['single-instance', 'delete-preserves-original']
            api('/api/shutdown', {})
            process.wait(timeout=30)
            assert process.returncode == 0
            print(json.dumps({'passed': True, 'system': platform.platform(), 'architecture': platform.machine(), 'checks': checks}))
        finally:
            if process.poll() is None:
                process.terminate()
                try:
                    process.wait(timeout=15)
                except subprocess.TimeoutExpired:
                    process.kill(); process.wait()


if __name__ == '__main__':
    main()
